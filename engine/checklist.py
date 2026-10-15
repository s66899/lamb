# -*- coding: utf-8 -*-
"""加工后核对 / 新购件核对 —— 数据层(本地端和网页端共用)。

本人要的是:「加工玩的零件的核对以及备注,和新购买零件的核对和备注」。
2026-10-15 又补了一句关键的话:

    **「采购件的核对是核对数量和型号」**

所以两栏**不是同一回事**,不能共用一套字段:

    加工后核对(自产件)  问的是「做出来没有、做得对不对」
        状态:待核对 / 合格 / 不合格 / 返工  + 备注

    新购件核对(外购件)  问的是「到货数量对不对、型号对不对」
        型号:待核对 / 正确 / 不对           + 到货数量 + 备注
        到货数量会自动跟需求量比,少了会标出差额 —— 外购件最常出事
        的就是「型号拿错」和「数量少几个」,让本人去点「合格/不合格」
        等于把这两件事都盖掉了。

三个决定(2026-10-13 与本人确认):
  1. 数据放**单独一个 `2_BOM表/核对表.xlsx`**,不给 BOM 表加列。
  2. **按分类自动分**:自产件看加工后核对,外购件看新购件核对。
  3. 两段式写入:先预览列清单,确认后才落盘。

★ 主键是 **(表名, 行号)**,不是图号:

    外购件里 22 行图号都叫「杯头螺钉」(M8*12 / M2*8 / M2*6 / M4*10 …),
    另有 14 行图号是空的。89 件只有 53 个不同的图号。
    按图号做主键,改一个 M8*12 标「合格」,另外 21 个在看板上
    **同时**变合格,而且不报错。图号为空的那 14 行会全部互相覆盖。
    `(表名, 行号)` 则 89 件全不重。

    行号会漂(BOM 表重建后顺序变),所以**图号是第二道校验**:
    写的时候对不上就整条拒绝并告诉本人是哪一行。

铁律(和 engine/edit.py 一致):
  1. 只写 `2_BOM表/核对表.xlsx`。**绝不碰** `2_BOM表/<分类>物料表.xlsx`。
  2. 先备份再写,原子写。
  3. 读不到 / 表坏了都只是「看不到核对结果」,**绝不让它把看板读挂**。

表结构(第一行标题,第二行表头,第三行起数据):

    表 | 行 | 分类 | 图号 | 名称
    加工后核对 | 加工后备注            ← 自产件用
    新购件型号 | 到货数量 | 新购件备注  ← 外购件用
"""
from __future__ import annotations

import os
import shutil
import time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

FILE_NAME = "核对表.xlsx"
TITLE = "加工后核对 / 新购件核对"
HDR = ["表", "行", "分类", "图号", "名称",
       "加工后核对", "加工后备注",
       "新购件型号", "到货数量", "新购件备注"]
COL_SRC, COL_ROW, COL_CAT, COL_PN, COL_NM = 1, 2, 3, 4, 5
COL_M_ST, COL_M_NOTE = 6, 7
COL_P_ST, COL_P_QTY, COL_P_NOTE = 8, 9, 10
HEADER_ROW = 2
FIRST_DATA_ROW = 3
N_COL = len(HDR)

MACHINED_CATS = ("机加工件", "塑胶件", "钣金件", "3D打印件")
PURCHASE_CATS = ("外购件",)

# 自产件:做出来没有、做得对不对
M_STATUSES = [("todo", "待核对"), ("ok", "合格"),
              ("bad", "不合格"), ("redo", "返工")]
# 外购件:型号对不对(本人原话:「采购件的核对是核对数量和型号」)
P_STATUSES = [("todo", "待核对"), ("ok", "型号正确"), ("bad", "型号不对")]

STATUSES_BY_KIND = {"m": M_STATUSES, "p": P_STATUSES}
# 保留 STATUSES 这个名字给外部(前端元数据、测试)用 —— 它现在代表「默认那套」,
# 也就是自产件那套;外购件那套要去 STATUSES_BY_KIND["p"] 取。
STATUSES = M_STATUSES
ST_NAME = dict(M_STATUSES)
P_ST_NAME = dict(P_STATUSES)
LABELS = {"m": "加工后核对", "p": "新购件核对"}
LABEL_FIELDS = {"m": ("st", "note"), "p": ("st", "qty", "note")}
MAX_NOTE = 500
DEFAULT = {"st": "todo", "note": "", "qty": None}


class CheckError(Exception):
    """校验不过。消息直接给本人看。"""


# ---------------------------------------------------------------- 工具
def check_path(proj) -> Path:
    return Path(proj.bom) / FILE_NAME


def statuses_of(kind: str):
    return STATUSES_BY_KIND.get(kind, M_STATUSES)


def st_name(kind: str, code: str) -> str:
    return dict(statuses_of(kind)).get(code, code)


def _code_of(kind: str, v):
    """把「ok」和「合格」都归一成代码「ok」。不认得的返回 None。

    网页端下拉送的是代码,但人也会在核对表里直接敲「型号不对」——
    两种都收下,省得改个字段就报「状态不认识」。
    """
    s = _norm(v)
    if not s:
        return "todo"
    m = dict(statuses_of(kind))
    if s in m:
        return s
    for code, name in m.items():
        if s == name:
            return code
    return None


def empty_for(it) -> dict:
    return {"m": dict(DEFAULT), "p": dict(DEFAULT)}


def kind_of(it) -> str:
    cat = str(it.get("cat") or "")
    if cat in MACHINED_CATS:
        return "m"
    if cat in PURCHASE_CATS:
        return "p"
    return ""


def need_qty(it):
    """需求量。取不到就返回 None —— 宁可不算差额,也不能编一个数字出来。"""
    for f in ("qty", "need"):
        v = it.get(f)
        if v is None or v == "":
            continue
        try:
            n = float(str(v).strip())
        except (TypeError, ValueError):
            continue
        return int(n) if n == int(n) else n
    return None


def short_of(it, got) -> float:
    """差几个。`got` 为空(还没记到货数量)时返回 0 —— 没填不算缺。"""
    need = need_qty(it)
    if need is None or got is None:
        return 0
    try:
        d = float(need) - float(got)
    except (TypeError, ValueError):
        return 0
    return d if d > 0 else 0


def done(it) -> bool:
    k = kind_of(it)
    if not k:
        return True
    c = (it.get("chk") or {}).get(k) or DEFAULT
    return c.get("st") in ("ok", "bad")


def bad_of(it) -> bool:
    """算不算「有问题」。

    外购件两件事都算:型号不对,或者记了到货数量但**比需求量少** ——
    数量短缺在采购里同样是「不合格」,只看型号会把缺货漏过去。
    """
    k = kind_of(it)
    if not k:
        return False
    c = (it.get("chk") or {}).get(k) or DEFAULT
    if c.get("st") == "bad":
        return True
    return short_of(it, c.get("qty")) > 0


def get(it) -> dict:
    k = kind_of(it)
    c = (it.get("chk") or {})
    return dict(c.get(k) or DEFAULT) if k else dict(DEFAULT)


def key_of(it) -> tuple:
    return (str(it.get("src") or ""), _norm_row(it.get("xrow")))


def tally(items, kind: str) -> dict:
    t = d = b = s = 0
    for it in items:
        if kind_of(it) != kind:
            continue
        t += 1
        c = (it.get("chk") or {}).get(kind) or DEFAULT
        if c.get("st") in ("ok", "bad"):
            d += 1
        if c.get("st") == "bad":
            b += 1
        if short_of(it, c.get("qty")) > 0:
            s += 1
    return {"total": t, "done": d, "bad": b, "todo": t - d, "short": s}


def _norm(v) -> str:
    return "" if v is None else str(v).strip()


def _norm_row(v):
    try:
        return str(int(v))
    except (TypeError, ValueError):
        return _norm(v)


def _norm_qty(v):
    """到货数量。可以是空(还没记),否则必须是数字。

    存成字符串是有意的 —— 需求量本来就可能是「2+2+2=6」这种写法,
    强行转 int 会在读取时就炸,而且保不住原样。
    """
    s = _norm(v)
    if not s:
        return None
    try:
        float(s)
    except ValueError:
        raise CheckError(f"到货数量「{s}」得是数字,或者留空(还没记到货)")
    return s


# ---------------------------------------------------------------- 读
def read(proj) -> dict:
    out = {}
    p = check_path(proj)
    if not p.is_file():
        return out
    try:
        wb = load_workbook(p, data_only=True, read_only=True)
    except Exception:
        return out
    try:
        ws = wb.active
        for row in ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
            if not row or row[COL_SRC - 1] is None:
                continue
            key = (_norm(row[COL_SRC - 1]), _norm_row(row[COL_ROW - 1]))
            if key[0] == "" or key[1] == "":
                continue          # 定位不了主键的行直接跳过,不做猜测
            rec = {"pn": _norm(row[COL_PN - 1]),
                   "cat": _norm(row[COL_CAT - 1]),
                   "m": {"st": _norm(row[COL_M_ST - 1]) or "todo",
                         "note": _norm(row[COL_M_NOTE - 1])},
                   "p": {"st": _norm(row[COL_P_ST - 1]) or "todo",
                         "note": _norm(row[COL_P_NOTE - 1]),
                         "qty": _norm(row[COL_P_QTY - 1]) or None}}
            for side in ("m", "p"):
                # ★ 这里存回去的必须是**代码**(ok/bad/todo),不是显示名。
                #   早先写反了:校验通过之后把 `names` 赋给了 st,
                #   于是 read() 出来的 st 是「合格」而不是「ok」——
                #   get()/done()/tally() 全部对不上,界面显示正常、
                #   统计永远是 0,还不报错。
                if _code_of(side, rec[side]["st"]) is None:
                    rec[side]["st"] = "todo"
                else:
                    rec[side]["st"] = _code_of(side, rec[side]["st"])
                rec[side]["note"] = rec[side]["note"][:MAX_NOTE]
            out[key] = rec
    except Exception:
        return out
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return out


def attach(cfg, proj, items) -> int:
    data = read(proj)
    n = 0
    for it in items:
        rec = data.get(key_of(it))
        if rec and rec["pn"] == _norm(it.get("pn")):
            it["chk"] = {"m": dict(rec["m"]), "p": dict(rec["p"])}
            n += 1
        else:
            it["chk"] = empty_for(it)
    return n


# ---------------------------------------------------------------- 写
def _norm_rec(cat, src, row, pn, nm, kind, st, note="", qty=None) -> dict:
    src, row = _norm(src), _norm_row(row)
    cat, pn = _norm(cat), _norm(pn)
    if not src or not row:
        raise CheckError("定位不到行(表名或行号是空的),写不进去 —— "
                         "这一件不是从物料表里读出来的?")
    if not cat:
        raise CheckError("分类是空的,写不进去")
    kind = _norm(kind)
    if kind not in LABELS:
        raise CheckError(f"核对栏 {kind!r} 不认识,只能是:"
                         + "、".join(LABELS[k] for k in LABELS))
    st = _norm(st) or "todo"
    code = _code_of(kind, st)
    if code is None:
        raise CheckError(
            f"「{LABELS[kind]}」的状态「{st}」不认识,只能是:"
            + "/".join(n for _, n in statuses_of(kind)))
    return {"cat": cat, "src": src, "row": row, "pn": pn, "nm": _norm(nm),
            "kind": kind, "st": code, "qty": _norm_qty(qty),
            "note": _norm(note)[:MAX_NOTE]}


def _backup(p: Path, log=print) -> str | None:
    if not p.is_file():
        return None
    d = p.parents[1] / "7_归档" / f"核对表改前备份_{time.strftime('%Y%m%d_%H%M%S')}"
    try:
        d.mkdir(parents=True, exist_ok=True)
        shutil.copy2(p, d / p.name)
        return str(d)
    except OSError as e:
        log(f"[备份失败] {e}")
        return None


def _atomic_save(wb, dest: Path):
    """同目录临时文件 -> os.replace。

    临时文件**必须和目标同一个目录**(跨盘符 os.replace 报 WinError 17),
    而且必须带 .xlsx 后缀(Excel 只认后缀,openpyxl 不挑)。
    """
    tmp = dest.with_name(dest.stem + f".tmp{os.getpid()}.xlsx")
    try:
        wb.save(str(tmp))
        os.replace(str(tmp), str(dest))
    finally:
        if tmp.exists():
            try:
                tmp.unlink()
            except OSError:
                pass
        try:
            wb.close()
        except Exception:
            pass


def _new_wb():
    wb = Workbook()
    ws = wb.active
    ws.title = "核对"
    ws.cell(1, 1, TITLE).font = Font(bold=True, size=12)
    ws.cell(1, COL_M_ST, "← 自产件用这两列").font = Font(size=10, color="FF808080")
    ws.cell(1, COL_P_ST, "← 外购件用这三列").font = Font(size=10, color="FF808080")
    fill = PatternFill("solid", fgColor="FF1F3B5C")
    fill2 = PatternFill("solid", fgColor="FF3B5C1F")
    for c, h in enumerate(HDR, 1):
        cell = ws.cell(HEADER_ROW, c, h)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = fill2 if c >= COL_P_ST else fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for c, w in enumerate((22, 7, 12, 20, 30, 12, 40, 12, 10, 40), 1):
        ws.column_dimensions[ws.cell(HEADER_ROW, c).column_letter].width = w
    ws.freeze_panes = ws.cell(FIRST_DATA_ROW, 1).coordinate
    return wb


def set_one(proj, rec: dict, log=print) -> dict:
    """写一条。按 (表名, 行号) 定位,图号做第二道校验;没有就加一行。"""
    p = check_path(proj)
    p.parent.mkdir(parents=True, exist_ok=True)
    new = not p.is_file()
    wb = _new_wb() if new else load_workbook(p)
    try:
        ws = wb.active
        # 表头对不上(被人改过 / 老版本)—— **停,不写**。
        # 早先悄悄建一张空表接着写,看着「成功了」,实际上把整份核对记录清空。
        if not new and _norm(ws.cell(HEADER_ROW, COL_PN).value) != "图号":
            wb.close()
            raise CheckError(
                f"{p.name} 的表头对不上(第 {HEADER_ROW} 行第 {COL_PN} 列"
                f"应该是「图号」,现在是"
                f"「{_norm(ws.cell(HEADER_ROW, COL_PN).value)}」)—— "
                f"可能是这张表被人手工改过,或者是很老的版本。"
                f"**没有写入**,原数据一点没动。"
                f"把它挪走或改名(比如加个 .bak 后缀),程序会自己建一张新的。")
        hit = None
        if not new:
            for r in range(FIRST_DATA_ROW, (ws.max_row or 2) + 1):
                if (_norm(ws.cell(r, COL_SRC).value) == rec["src"]
                        and _norm_row(ws.cell(r, COL_ROW).value) == rec["row"]):
                    hit = r
                    break
        if hit is not None:
            onfile = _norm(ws.cell(hit, COL_PN).value)
            if onfile and rec["pn"] and onfile != rec["pn"]:
                raise CheckError(
                    f"{rec['src']} 第 {rec['row']} 行,核对表里记的是图号"
                    f"「{onfile}」,现在这一行是「{rec['pn']}」——"
                    f"表重建过了,行号对不上。为免写到别的零件上,"
                    f"这一条没有写入。")
        else:
            hit = max((ws.max_row or HEADER_ROW) + 1, FIRST_DATA_ROW)
            ws.cell(hit, COL_SRC, rec["src"])
            ws.cell(hit, COL_ROW, int(rec["row"]) if rec["row"].isdigit()
                    else rec["row"])
            ws.cell(hit, COL_CAT, rec["cat"])
            ws.cell(hit, COL_PN, rec["pn"])
            ws.cell(hit, COL_NM, rec["nm"])
        if rec["kind"] == "m":
            ws.cell(hit, COL_M_ST, rec["st"])
            ws.cell(hit, COL_M_NOTE, rec["note"])
        else:
            ws.cell(hit, COL_P_ST, rec["st"])
            ws.cell(hit, COL_P_QTY, rec["qty"])
            ws.cell(hit, COL_P_NOTE, rec["note"])
        bak = _backup(p, log)
        _atomic_save(wb, p)
    except Exception:
        try:
            wb.close()
        except Exception:
            pass
        raise
    extra = ""
    if rec["kind"] == "p" and rec["qty"]:
        extra = f" 到货 {rec['qty']}"
    log(f"  核对表 {rec['src']} 第 {rec['row']} 行 {rec['pn']} "
        f"{LABELS[rec['kind']]} -> {st_name(rec['kind'], rec['st'])}{extra}"
        + (f" 备注「{rec['note'][:20]}」" if rec["note"] else ""))
    return {"ok": True, "file": str(p), "backup": bak, "row": hit,
            "created": new}


def set(cfg, proj, it: dict, kind: str, code: str, note: str = "",
        qty=None, log=print) -> bool:
    """本地端 UI 的**唯一**入口:把一件的核对结果写盘。

    为什么不直接用 `set_one()`:那个要 (src, row, pn) 三样,
    而 UI 手上只有一个 item —— 早先 `_set_chk` 顺手写了
    `checklist.set(...)`,而这个函数**根本不存在**,于是 AttributeError
    被 messagebox 吃掉,界面上「点了一下没反应」。
    入口收在这一处,以后加校验、加备份策略都只改一个地方。
    """
    try:
        set_one(proj, _norm_rec(it.get("cat"), it.get("src"), it.get("xrow"),
                                it.get("pn"), it.get("nm"), kind, code,
                                note=note, qty=qty), log=log)
        return True
    except CheckError:
        raise
    except Exception as e:
        raise CheckError(f"{type(e).__name__}: {e}")


def _plan(cfg, proj, edits) -> dict:
    """算出这一批改动真正要落盘的东西(内部用)。"""
    if not isinstance(edits, list) or not edits:
        raise CheckError("没有要保存的修改")
    if len(edits) > 500:
        raise CheckError(f"一次最多 500 条,现在有 {len(edits)} 条")
    cur = read(proj)
    rows = []
    for i, e in enumerate(edits, 1):
        if not isinstance(e, dict):
            raise CheckError(f"第 {i} 条格式不对")
        rec = _norm_rec(e.get("cat"), e.get("src"), e.get("row"),
                        e.get("pn"), e.get("nm"), e.get("kind"),
                        e.get("st"), note=e.get("note"), qty=e.get("qty"))
        old = (cur.get((rec["src"], rec["row"])) or {}).get(rec["kind"]) or DEFAULT
        if (old["st"] == rec["st"] and old["note"] == rec["note"]
                and old.get("qty") == rec["qty"]):
            continue                      # 没真变的不列
        was = [st_name(rec["kind"], old["st"])]
        if old.get("qty"):
            was.append("到货 " + str(old["qty"]))
        if old["note"]:
            was.append(old["note"][:30])
        now = [st_name(rec["kind"], rec["st"])]
        if rec["kind"] == "p" and rec["qty"]:
            now.append("到货 " + str(rec["qty"]))
        if rec["note"]:
            now.append(rec["note"][:30])
        rows.append({"cat": rec["cat"], "src": rec["src"], "row": rec["row"],
                     "pn": rec["pn"], "nm": rec["nm"],
                     "kind": rec["kind"], "label": LABELS[rec["kind"]],
                     "was": "、".join(was), "now": "、".join(now),
                     "_rec": rec})
    return {"ok": True, "rows": rows, "total": len(rows)}


def _public(plan: dict) -> dict:
    return {"ok": True, "total": plan["total"],
            "rows": [{k: v for k, v in r.items() if not k.startswith("_")}
                     for r in plan["rows"]]}


def preview(cfg, proj, edits) -> dict:
    """和 engine/edit.py 一样两段式:先给本人看清单,确认后才写。"""
    return _public(_plan(cfg, proj, edits))


def apply(cfg, proj, edits, log=print) -> dict:
    plan = _plan(cfg, proj, edits)
    bak = None
    for r in plan["rows"]:
        res = set_one(proj, r["_rec"], log=log)
        bak = bak or res.get("backup")
    return {"ok": True, "written": plan["total"], "backup": bak,
            "rows": _public(plan)["rows"]}